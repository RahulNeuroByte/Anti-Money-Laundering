import json
from pathlib import Path

import pytest
from docx import Document

from aml.report import build_report

ROOT = Path(__file__).resolve().parents[1]
REQUIRED = [
    "Abstract and problem statement",
    "Dataset description and preprocessing",
    "Exploratory data analysis",
    "Feature engineering, imbalance handling and validation strategy",
    "Model building, tuning, evaluation and selection",
    "Model serialisation (.pkl) and validation",
    "FastAPI design, testing and CI pipeline",
    "Azure MLflow implementation",
    "Results, limitations, ethical considerations and future work",
    "References",
]


def _text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_report_builds_without_inputs_and_never_invents_results(tmp_path):
    out = build_report(tmp_path, tmp_path / "r.docx")
    text = _text(out)
    for heading in REQUIRED:
        assert heading in text
    assert text.count("[SCREENSHOT PLACEHOLDER]") >= 9
    assert "Not available" in text
    assert "Deployment status: [fill in" in text


def test_report_uses_values_from_result_files(tmp_path):
    (tmp_path / "dataset_profile.json").write_text(
        json.dumps(
            {
                "shape": [10, 3],
                "memory_mb": 1.0,
                "dtypes": {"Amount": "float32"},
                "missing": {"Amount": 0},
                "duplicate_rows": 0,
                "amount_summary": {"count": 10.0, "mean": 5.0},
                "categorical_cardinality": {},
                "unique_senders": 4,
                "unique_receivers": 5,
                "timestamp_range": ["2023-01-01", "2023-01-02"],
                "target": {
                    "name": "Is_laundering",
                    "positives": 1,
                    "negatives": 9,
                    "positive_rate_percent": 10.0,
                    "imbalance_ratio_neg_to_pos": 9.0,
                },
            }
        )
    )
    text = _text(build_report(tmp_path, tmp_path / "r.docx"))
    assert "10 transactions" in text and "Imbalance ratio" in text


@pytest.mark.skipif(not (ROOT / "reports" / "model_results.json").exists(), reason="no results")
def test_report_with_real_results_names_selected_model(tmp_path):
    results = json.loads((ROOT / "reports" / "model_results.json").read_text())
    text = _text(build_report(ROOT / "reports", tmp_path / "r.docx"))
    assert results["selected"] in ("logreg", "tree", "forest")
    assert f"{results[results['selected']]['test']['pr_auc']:.3f}" in text
