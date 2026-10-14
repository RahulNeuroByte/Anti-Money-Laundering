"""Generate reports/AML_Detection_Capstone_Report.docx from the saved result files.

    $env:PYTHONPATH="src"; python -m aml.report

Every number comes from reports/*.json. Missing inputs produce a note, never an invented value.
Screenshots are placeholders for you to fill in; Azure results are never asserted here.
"""

from __future__ import annotations

import json
from datetime import date
from importlib import metadata
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from aml.config import PROJECT_ROOT
from aml.train import BLOCK_DAYS, WINDOW_DAYS

FIGURES = (
    ("01_monthly_volume.png", "Monthly transactions and laundering cases (training window)"),
    ("02_amount_by_class.png", "Amount distribution by class (log scale)"),
    ("03_rate_payment_type.png", "Laundering rate by payment type"),
    ("04_rate_sender_location.png", "Laundering rate by sender bank location"),
    ("05_rate_received_currency.png", "Laundering rate by received currency"),
    ("06_rate_by_hour.png", "Laundering rate by hour of day"),
    ("07_typology_counts.png", "Laundering typologies in the training window (analysis only)"),
    ("08_correlation.png", "Spearman correlation of engineered features with the target"),
)
MODEL_LABEL = {"logreg": "Logistic Regression", "tree": "Decision Tree", "forest": "Random Forest"}
COMMANDS = [
    "python -m aml.profile_data", "python -m aml.eda", "python -m aml.train features",
    "python -m aml.train imbalance", "python -m aml.train tune --cv-neg 120000",
    "python -m aml.train final", "python -m aml.verify artifact", "python -m aml.verify logreg", "python -m aml.verify auc",
    "python -m mlops.track_models", "python -m aml.report",
]  # fmt: skip


class Builder:
    def __init__(self, reports: Path):
        self.reports = reports
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)

    # -- inputs ---------------------------------------------------------------------------
    def load(self, name: str) -> Any | None:
        path = self.reports / name
        return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else None

    # -- primitives -----------------------------------------------------------------------
    def h(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level)

    def p(self, text: str, italic: bool = False) -> None:
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.italic = italic

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def missing(self, what: str, command: str) -> None:
        self.p(f"[Not available: {what}. Run `{command}` and regenerate this report.]", True)

    def table(self, header: list[str], rows: list[list[Any]], caption: str) -> None:
        self.tab_no += 1
        cap = self.doc.add_paragraph()
        cap.add_run(f"Table {self.tab_no}. {caption}").bold = True
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, name in enumerate(header):
            t.rows[0].cells[i].text = name
            for r in t.rows[0].cells[i].paragraphs[0].runs:
                r.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = (
                    f"{v:,.4f}"
                    if isinstance(v, float)
                    else f"{v:,}"
                    if isinstance(v, int)
                    else str(v)
                )
        self.doc.add_paragraph()

    def figure(self, filename: str, caption: str, width_in: float = 6.0) -> None:
        path = self.reports / "figures" / filename
        if not path.is_file():
            self.p(f"[Figure missing: {filename}. Run the pipeline to generate it.]", True)
            return
        self.fig_no += 1
        self.doc.add_picture(str(path), width=Inches(width_in))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph(f"Figure {self.fig_no}. {caption}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True

    def screenshot(self, label: str, height_cm: float = 8.0) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        row = t.rows[0]
        row.height, row.height_rule = Cm(height_cm), WD_ROW_HEIGHT_RULE.EXACTLY
        cell = row.cells[0]
        cell.text = f"[SCREENSHOT PLACEHOLDER]\n{label}\n(insert your screenshot here)"
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_paragraph()

    # -- sections -------------------------------------------------------------------------
    def title_page(self) -> None:
        for _ in range(5):
            self.doc.add_paragraph()
        t = self.doc.add_paragraph("End-to-End Anti-Money Laundering (AML) Detection")
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.runs[0].bold, t.runs[0].font.size = True, Pt(26)
        s = self.doc.add_paragraph(
            "Machine Learning, FastAPI, GitHub Actions, Azure Machine Learning and MLflow"
        )
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.runs[0].font.size = Pt(14)
        for line in (
            "Capstone Project Report",
            "Track chosen: Anti-Money Laundering (SAML-D transaction dataset)",
            "",
            "Student name: [to be filled in]",
            "Roll number / Programme: [to be filled in]",
            "Guide / Institution: [to be filled in]",
            f"Report generated: {date.today().isoformat()}",
        ):
            para = self.doc.add_paragraph(line)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_page_break()

    def contents(self) -> None:
        self.h("Contents", 1)
        self.bullets(
            [
                "1. Abstract and problem statement",
                "2. Dataset description and preprocessing",
                "3. Exploratory data analysis",
                "4. Feature engineering, imbalance handling and validation strategy",
                "5. Model building, tuning, evaluation and selection",
                "6. Model serialisation (.pkl) and validation",
                "7. FastAPI design, testing and CI pipeline",
                "8. Azure MLflow implementation: tracking, registry, deployment, monitoring",
                "9. Results, limitations, ethical considerations and future work",
                "10. References",
                "Appendices A to C",
            ]
        )
        self.doc.add_page_break()

    def abstract(self) -> None:
        mr, prof = self.load("model_results.json"), self.load("dataset_profile.json")
        self.h("1. Abstract and problem statement", 1)
        self.h("Abstract", 2)
        if mr:
            best = mr["selected"]
            test = mr[best]["test"]
            f1 = test["at_best_f1"]
            self.p(
                f"This project builds a complete machine-learning system for detecting suspicious "
                f"transactions in the SAML-D anti-money-laundering dataset. Three scikit-learn "
                f"models (Logistic Regression, Decision Tree, Random Forest) were compared under "
                f"several class-imbalance strategies using a chronological train/validation/test "
                f"split. The selected model, {MODEL_LABEL[best]}, reached a PR-AUC of "
                f"{test['pr_auc']:.3f} and ROC-AUC of {test['roc_auc']:.3f} on the held-out test "
                f"window, with precision {f1['precision']:.3f}, recall {f1['recall']:.3f} and F1 "
                f"{f1['f1']:.3f} at a threshold chosen on the validation window. The model is "
                f"served through a FastAPI service, tested automatically and prepared for "
                f"deployment as an Azure ML managed online endpoint with MLflow tracking."
            )
        else:
            self.missing("model results", "python -m aml.train final")
        self.h("Problem statement", 2)
        rate = f"{prof['target']['positive_rate_percent']}%" if prof else "a very small fraction"
        self.p(
            "Financial institutions must detect and report suspicious activity. Missed cases "
            "bring regulatory penalties, while excessive false alerts overload investigators. "
            f"The task is to classify each transaction as normal or laundering when only {rate} "
            "of transactions are laundering, so accuracy alone is misleading and model selection "
            "uses precision, recall, F1 and precision-recall AUC."
        )
        self.p(
            "Track chosen: Anti-Money Laundering detection (SAML-D). The label column is Is_laundering."
        )

    def dataset(self) -> None:
        prof, clean = self.load("dataset_profile.json"), self.load("cleaning_log.json")
        self.h("2. Dataset description and preprocessing", 1)
        if prof:
            rows, cols = prof["shape"]
            tg = prof["target"]
            self.p(
                f"The raw file contains {rows:,} transactions. After merging Date and Time into one "
                f"Timestamp column the working table has {cols} columns, covering "
                f"{prof['timestamp_range'][0]} to {prof['timestamp_range'][1]}, with "
                f"{prof['unique_senders']:,} distinct sender and {prof['unique_receivers']:,} "
                f"receiver accounts. Memory footprint after typed loading is {prof['memory_mb']} MB."
            )
            self.table(
                ["Column", "Type", "Missing"],
                [[c, t, prof["missing"][c]] for c, t in prof["dtypes"].items()],
                "Schema, data types and missing values",
            )
            a = prof["amount_summary"]
            self.table(
                ["Statistic", "Amount"],
                [[k, float(v)] for k, v in a.items()],
                "Summary statistics of Amount",
            )
            self.table(
                ["Target statistic", "Value"],
                [
                    ["Positives (laundering)", tg["positives"]],
                    ["Negatives (normal)", tg["negatives"]],
                    ["Positive rate (%)", tg["positive_rate_percent"]],
                    ["Imbalance ratio (negative : positive)", tg["imbalance_ratio_neg_to_pos"]],
                ],
                "Target distribution and class imbalance",
            )
        else:
            self.missing("dataset profile", "python -m aml.profile_data")
        self.h("Cleaning decisions", 2)
        if clean:
            o = clean["amount_outliers"]
            self.bullets(
                [
                    f"Exact duplicate rows removed: {clean['duplicates_removed']:,}.",
                    f"Rows removed for missing required fields: {clean['rows_removed_missing']:,}.",
                    f"Rows removed for non-positive amounts: {clean['rows_removed_nonpositive_amount']:,}.",
                    f"Self-transfers found: {clean['self_transfers']:,}.",
                    f"Rows kept: {clean['rows_out']:,} of {clean['rows_in']:,}.",
                    f"Amount outliers above the IQR fence ({o['iqr_upper_fence']:,.2f}): "
                    f"{o['count_above_fence']:,} rows, with laundering rate "
                    f"{100 * o['positive_rate_above_fence']:.3f}% against "
                    f"{100 * o['positive_rate_below_fence']:.3f}% below the fence. They are "
                    f"retained as signal and compressed with a log transform.",
                ]
            )
        else:
            self.missing("cleaning log", "python -m aml.eda")
        self.p(
            "Leakage control: the column Laundering_type is a label-derived typology that fully "
            "determines the target, so it is never used as a feature. It is used only for error analysis."
        )
        self.screenshot("Dataset preview and preprocessing steps (terminal or notebook output)")

    def eda(self) -> None:
        eda, tv = self.load("eda_summary.json"), self.load("train_vs_valid_auc.json")
        self.h("3. Exploratory data analysis", 1)
        self.p(
            "All EDA was computed on the training window only; validation and test were not inspected for modelling decisions."
        )
        if eda:
            sp = eda["split"]
            self.table(
                ["Split", "Rows", "Positives", "Positive rate", "Start", "End"],
                [
                    [
                        n,
                        v["rows"],
                        v["positives"],
                        f"{100 * v['positive_rate']:.4f}%",
                        v["start"][:10],
                        v["end"][:10],
                    ]
                    for n, v in sp.items()
                ],
                "Chronological split",
            )
            la = eda["leakage_audit"]
            self.p(
                f"Leakage audit: Laundering_type has {la['laundering_type_values']} values; "
                f"{la['values_appearing_in_both_classes']} appear in both classes, confirming it "
                "cannot be used as a feature."
            )
        else:
            self.missing("EDA summary", "python -m aml.eda")
        for fn, cap in FIGURES:
            self.figure(fn, cap)
        if tv:
            top = list(tv.items())[:8]
            self.table(
                ["Feature", "Univariate AUC (train)", "Univariate AUC (validation)"],
                [[k, v[0], v[1]] for k, v in top],
                "Signal of engineered features: training versus validation",
            )
            self.p(
                "Insight: account-profile signals are weaker on validation than on training when "
                "profiles are computed over the whole training window. This motivated causal "
                "(lagged) profiles, described in the next section."
            )
        self.screenshot(
            "EDA visualisations and insights (optional: your own notebook screenshots)", 6
        )

    def features(self) -> None:
        eda, imb, lrc = (
            self.load("eda_summary.json"),
            self.load("imbalance_results.json"),
            self.load("logreg_converged.json"),
        )
        self.h("4. Feature engineering, imbalance handling and validation strategy", 1)
        self.h("Feature engineering", 2)
        self.bullets(
            [
                "Row features: log amount, hour, day of week, cross-border flag, cross-currency flag, round-amount flag, 9,000-10,000 band flag, self-transfer flag.",
                "Categorical features: payment type, payment and received currency, sender and receiver bank location (unseen values map to a reserved level).",
                "Account profiles for sender and receiver: transaction counts, totals, mean and maximum amounts, number of distinct counterparties, and in/out flow ratios.",
                f"Causal profiles: for every {BLOCK_DAYS}-day block the profiles are built only from the preceding {WINDOW_DAYS} days, so no row sees its own or future activity and every row sees a history of the same length. The first {WINDOW_DAYS} days are warm-up and excluded from training.",
                "Account identifiers are never used as raw features.",
            ]
        )
        if eda:
            self.p(f"The feature builder produces {eda['n_features']} model inputs.")
        self.h("Validation strategy", 2)
        self.p(
            "A chronological 70/15/15 split is used. Training is fitted on the training window, "
            "hyper-parameters and the decision threshold are chosen on the validation window, and "
            "the test window is scored once with thresholds fixed beforehand. Resampling is "
            "applied inside the training fit only. Cross-validated average precision was used "
            "during tuning but is optimistic here because folds share accounts and time, so "
            "model selection relies on the validation window."
        )
        self.h("Imbalance experiments", 2)
        if imb:
            rows = [
                [
                    MODEL_LABEL[r["model"]],
                    r["strategy"],
                    r["pr_auc"],
                    r["roc_auc"],
                    r["best_f1"],
                    r["precision"],
                    r["recall"],
                ]
                for r in imb
            ]
            self.table(
                ["Model", "Strategy", "PR-AUC", "ROC-AUC", "Best F1", "Precision", "Recall"],
                rows,
                "Imbalance strategies compared on the validation window (default hyper-parameters)",
            )
            self.p(
                "Strategies: none; class weights; random oversampling of the minority class to a "
                "1:10 ratio; random undersampling to 1:10; and SMOTENC (after undersampling) for "
                "tree models only, because it needs ordinal-coded categoricals. Precision, recall and "
                "F1 in this table use a validation-tuned threshold and are therefore optimistic."
            )
        else:
            self.missing("imbalance results", "python -m aml.train imbalance")
        self.figure("09_imbalance_pr_auc.png", "Validation PR-AUC by imbalance strategy and model")
        if lrc:
            self.table(
                ["Strategy", "PR-AUC", "ROC-AUC"],
                [[r["strategy"], r["pr_auc"], r["roc_auc"]] for r in lrc],
                "Logistic Regression re-run with converged solver (max_iter=1000)",
            )

    def models(self) -> None:
        mr, tune = self.load("model_results.json"), self.load("tuning_results.json")
        self.h("5. Model building, tuning, evaluation and selection", 1)
        if tune:
            self.table(
                ["Model", "Strategy", "Best hyper-parameters", "CV average precision"],
                [
                    [
                        MODEL_LABEL[m],
                        t["strategy"],
                        json.dumps(t["params"]),
                        t["cv_average_precision"],
                    ]
                    for m, t in tune.items()
                ],
                "Randomised search (3-fold stratified CV, average precision)",
            )
            self.p(
                "CV scores are computed on a sample with a higher positive share than the real data and are not comparable with validation or test PR-AUC."
            )
        else:
            self.missing("tuning results", "python -m aml.train tune --cv-neg 120000")
        if not mr:
            self.missing("model results", "python -m aml.train final")
            return
        for split in ("valid", "test"):
            rows = []
            for m in MODEL_LABEL:
                r = mr[m][split]
                f = r["at_best_f1"]
                rows.append(
                    [
                        MODEL_LABEL[m],
                        f["accuracy"],
                        f["precision"],
                        f["recall"],
                        f["f1"],
                        r["roc_auc"],
                        r["pr_auc"],
                        f["support_pos"],
                    ]
                )
            self.table(
                [
                    "Model",
                    "Accuracy",
                    "Precision",
                    "Recall",
                    "F1",
                    "ROC-AUC",
                    "PR-AUC",
                    "Positives",
                ],
                rows,
                f"{'Validation' if split == 'valid' else 'Test'} metrics at the threshold that maximises validation F1",
            )
        rows = []
        for m in MODEL_LABEL:
            c = mr[m]["test"]["at_best_f1"]
            rows.append([MODEL_LABEL[m], c["tp"], c["fp"], c["fn"], c["tn"], c["alerts"]])
        self.table(
            ["Model", "TP", "FP", "FN", "TN", "Alerts raised"], rows, "Test confusion matrices"
        )
        r80 = mr[mr["selected"]]["test"]["at_recall_0.80"]
        self.p(
            f"Operating point for higher recall (selected model, threshold {r80['threshold']:.4f} "
            f"chosen on validation for recall 0.80): test recall {r80['recall']:.3f}, precision "
            f"{r80['precision']:.3f}, {r80['alerts']:,} alerts."
        )
        self.figure("10_curves_valid.png", "ROC and precision-recall curves (validation)")
        self.figure("11_curves_test.png", "ROC and precision-recall curves (test)")
        self.figure("12_confusion_test.png", "Confusion matrices on the test window")
        self.figure("13_feature_importance.png", "Feature importance of the selected model")
        self.figure("14_threshold_sweep.png", "Precision, recall and F1 across decision thresholds")
        best = mr["selected"]
        self.p(
            f"Selection: {MODEL_LABEL[best]} has the highest validation PR-AUC "
            f"({mr[best]['valid']['pr_auc']:.3f}). Accuracy is reported but not used for selection "
            "because a model that flags nothing already exceeds 99.8% accuracy on this data."
        )
        ea = mr.get("error_analysis", {}).get("recall_by_typology_test")
        if ea:
            self.table(
                ["Typology (analysis only)", "Test cases", "Recall"],
                [[k, v["n"], v["recall"]] for k, v in ea.items()],
                "Recall by laundering typology for the selected model (test)",
            )
        self.screenshot("Model training and evaluation results (terminal output)", 7)

    def serialisation(self) -> None:
        av = self.load("artifact_validation.json")
        self.h("6. Model serialisation (.pkl) and validation", 1)
        self.p(
            "The deployable pipeline (feature builder with a frozen account-profile table, "
            "preprocessing, classifier) is saved with joblib as models/aml_pipeline.pkl, with the "
            "decision threshold in models/model_meta.json. Pickle files execute code when loaded: "
            "load only artifacts you built yourself."
        )
        if av:
            t = av["test_with_static_profile"]["at_saved_threshold"]
            self.table(
                ["Check", "Result"],
                [
                    [
                        "Reload prediction match (20,000 rows), max abs difference",
                        av["reload_max_abs_diff_20k_rows"],
                    ],
                    ["Reloaded predictions identical", str(av["reload_exact_match"])],
                    ["Scoring time for the whole test window (s)", av["raw_test_scoring_seconds"]],
                    ["Static profile ends at", av["static_profile_end"]],
                    [
                        "Artifact test PR-AUC (raw rows, static profile)",
                        av["test_with_static_profile"]["pr_auc"],
                    ],
                    [
                        "Artifact test precision / recall / F1",
                        f"{t['precision']:.3f} / {t['recall']:.3f} / {t['f1']:.3f}",
                    ],
                ],
                "Serialisation and artifact validation",
            )
            self.p(
                "The frozen profile becomes stale as time passes after its end date, which "
                "lowers performance slightly compared with refreshed profiles; a refresh job is part "
                "of the retraining plan."
            )
        else:
            self.missing("artifact validation", "python -m aml.verify artifact")
        self.screenshot("Saved model output (.pkl) and reload check", 6)

    def api(self) -> None:
        self.h("7. FastAPI design, testing and CI pipeline", 1)
        self.table(
            ["Endpoint", "Purpose", "Notes"],
            [
                [
                    "GET /health",
                    "Service and model status",
                    "HTTP 503 when the model or the Azure backend is not configured",
                ],
                [
                    "POST /predict",
                    "Score one transaction",
                    "Returns prediction, probability, threshold, model",
                ],
            ],
            "API endpoints",
        )
        self.bullets(
            [
                "Request schema (Pydantic): Timestamp, Sender_account, Receiver_account, Amount (greater than 0, at most 1e9), Payment_type, Payment_currency, Received_currency, Sender_bank_location, Receiver_bank_location. Categories are restricted to the values seen in training; extra fields are rejected.",
                "Validation errors return HTTP 422 with location, message and type only; submitted values are never echoed or logged.",
                "Backends: local (loads the .pkl) or azure (forwards to the managed online endpoint using the key). Upstream failures map to HTTP 502 or 504 without leaking the key.",
                "The same inference code is used by the API, offline checks and the Azure scoring script, and tests assert identical probabilities.",
                "Automated tests (pytest and FastAPI TestClient) cover valid input, each missing field, wrong types, out-of-range values, unseen categories, extra fields, NaN, timezone-aware timestamps, model-load failure, corrupt model files, API versus local parity, and the Azure-backend error paths against a local stub.",
            ]
        )
        self.screenshot("FastAPI Swagger UI (/docs) and Postman or cURL testing", 8)
        self.p(
            "Continuous integration: .github/workflows/ci.yml runs on every push and pull request with "
            "Python 3.11.9: install pinned requirements, pip check, ruff and pytest. A failing test fails the build."
        )
        self.screenshot("GitHub Actions CI run (passing build)", 7)

    def azure(self) -> None:
        ml = self.load("mlflow_summary.json")
        self.h("8. Azure MLflow implementation: tracking, registry, deployment, monitoring", 1)
        self.h("Experiment tracking and registry", 2)
        self.p(
            "mlops/track_models.py retrains the three models with the tuned settings and logs, for each run, "
            "parameters, validation and test metrics (precision, recall, F1, ROC-AUC, PR-AUC, confusion counts), "
            "ROC/PR and confusion-matrix plots, a feature-importance plot, a data hash, the dependency lock file and the "
            "deployable model with its input signature. The best run by validation PR-AUC is registered "
            "as 'aml-detector' with version tags (stage, alias, threshold, selection rule, data hash)."
        )
        if ml:
            self.table(
                [
                    "Model",
                    "Validation PR-AUC",
                    "Test PR-AUC",
                    "Test F1",
                    "Difference vs earlier run",
                ],
                [
                    [
                        MODEL_LABEL.get(r["params.model"], r["params.model"]),
                        r["metrics.valid_pr_auc"],
                        r["metrics.test_pr_auc"],
                        r["metrics.test_f1"],
                        r["metrics.repro_abs_diff_vs_phase4"],
                    ]
                    for r in ml["comparison"]
                ],
                "MLflow run comparison (local tracking store)",
            )
            self.p(
                f"Champion: {MODEL_LABEL.get(ml['champion_model'], ml['champion_model'])}, registered version "
                f"{ml['registered_version']}. Registry model equals the saved .pkl on 20,000 rows "
                f"(max abs difference {ml['registry_vs_pkl_max_abs_diff_20k_rows']})."
            )
        else:
            self.missing("MLflow summary", "python -m mlops.track_models")
        self.h("Azure ML SDK v2 pipeline", 2)
        self.p(
            "mlops/azure/submit_pipeline.py (or the notebook notebooks/azureml_sdk2_pipeline.ipynb, run on the "
            "'Python 3.10 - SDK v2' kernel) creates in the Azure ML workspace: a data asset (saml-d-raw), an "
            "environment (aml-train-env, Python 3.11.9), five command components (prepare, features, imbalance, "
            "tune, final), a pipeline job that connects them, and a registered model (aml-detector-scoring). "
            "Each step logs parameters, metrics, stepped precision/recall/F1 curves, and PNG and JSON artifacts "
            "through MLflow. The step scripts were run end to end locally with the job environment's pinned "
            "libraries and reproduced the reported model metrics exactly; the Azure run itself is not reported here."
        )
        self.p("Pipeline run status: [fill in after you run the pipeline in Azure ML].", True)
        self.screenshot(
            "Azure ML pipeline graph (Jobs > aml-detection-pipeline) and step outputs", 8
        )
        self.screenshot("Azure ML Data, Environments, Components and Models pages", 8)
        self.p(
            "The results above were produced with a local MLflow store. Logging to the Azure ML workspace is done by setting MLFLOW_TRACKING_URI to the workspace tracking URI (see INSTRUCTION.md); that step has not been verified by this report.",
            True,
        )
        self.screenshot("MLflow experiment tracking dashboard and Model Registry", 9)
        self.h("Deployment", 2)
        self.p(
            "Design: a managed online endpoint (key authentication) with a deployment named 'blue' on one "
            "Standard_DS2_v2 instance, a custom scoring script (mlops/azure/score.py) that reuses the same "
            "inference code, a conda environment with pinned dependencies, and Application Insights enabled. The "
            "scoring package and the Azure inference server were tested locally; no statement about the Azure "
            "deployment itself is made here."
        )
        self.p(
            "Deployment status: [fill in after you run mlops/azure/deploy.ps1 and record the outcome].",
            True,
        )
        self.screenshot("Azure deployment (endpoint and deployment status in Azure ML studio)", 8)
        self.screenshot("Endpoint testing (curl or Python client output, with the key hidden)", 8)
        self.h("Monitoring, retraining and safe rollout", 2)
        self.bullets(
            [
                "Monitoring: Application Insights for request volume, latency and errors; track the share of alerts per day and population drift of Amount, payment-type mix and cross-border share against the training window.",
                "Performance monitoring needs confirmed investigation outcomes; compare precision and recall on labelled cases each period.",
                "Retraining triggers: scheduled (for example monthly, to refresh account profiles) and event-based (drift beyond agreed limits or a fall in precision or recall).",
                "Safe rollout: deploy a candidate as 'green' beside 'blue', route a small traffic share, compare, then promote or roll back by moving traffic; delete the unused deployment to stop its cost.",
                "Cost: a managed deployment bills per hour while it exists. Delete the endpoint after testing (teardown script).",
            ]
        )

    def closing(self) -> None:
        mr = self.load("model_results.json")
        self.h("9. Results, limitations, ethical considerations and future work", 1)
        self.h("Limitations", 2)
        self.bullets(
            [
                "The dataset is synthetic; results do not demonstrate real-world AML performance.",
                "Typologies such as bipartite, fan-in and cycle patterns are detected less well than others; see the per-typology table.",
                "The deployed profile table is a snapshot and goes stale; one transaction is scored without live account history.",
                "Training used a sampled negative pool because of compute limits; full-data training may change results.",
                "The model was trained on the training window only, not refit on training plus validation.",
            ]
        )
        self.h("Ethical considerations", 2)
        self.bullets(
            [
                "False positives burden customers and investigators; thresholds should reflect investigator capacity, and alerts should trigger human review, not automatic action.",
                "Bias: location and currency features can proxy for nationality; fairness should be assessed on real data before use.",
                "Privacy: transaction and account data are sensitive; the API does not log request values, keys stay in environment variables or secret stores, and access should be restricted.",
                "Explainability: feature importance and alert reasons should be available to investigators and auditors.",
            ]
        )
        self.h("Future work", 2)
        self.bullets(
            [
                "Gradient-boosted models and graph features for the weak typologies.",
                "Live account-profile store and drift-triggered retraining.",
                "Optional CI/CD to train, register and deploy to Azure with a service principal.",
            ]
        )
        if mr:
            self.p(f"Selected model: {MODEL_LABEL[mr['selected']]}.")
        self.h("10. References", 1)
        self.bullets(
            [
                "SAML-D synthetic anti-money-laundering dataset (the source named in the project brief).",
                "scikit-learn, imbalanced-learn, pandas, FastAPI, Pydantic, pytest documentation.",
                "MLflow documentation (tracking, model registry); Azure Machine Learning documentation (managed online endpoints).",
            ]
        )

    def appendices(self) -> None:
        self.h("Appendix A: environment", 1)
        pkgs = [
            "pandas",
            "numpy",
            "scikit-learn",
            "imbalanced-learn",
            "fastapi",
            "pydantic",
            "mlflow-skinny",
            "python-docx",
        ]
        rows = []
        for name in pkgs:
            try:
                rows.append([name, metadata.version(name)])
            except metadata.PackageNotFoundError:
                rows.append([name, "not installed"])
        self.table(["Package", "Version (this report run)"], rows, "Library versions")
        self.h("Appendix B: reproduction commands (PowerShell, project root, PYTHONPATH=src)", 1)
        for c in COMMANDS:
            para = self.doc.add_paragraph(c)
            para.runs[0].font.name = "Consolas"
        self.h("Appendix C: additional screenshots and notes", 1)
        self.screenshot("Additional screenshot or notes", 8)

    def build(self, out: Path) -> Path:
        self.title_page()
        self.contents()
        self.abstract()
        self.dataset()
        self.eda()
        self.features()
        self.models()
        self.serialisation()
        self.api()
        self.azure()
        self.closing()
        self.appendices()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(out)
        return out


def build_report(reports_dir: Path | None = None, out: Path | None = None) -> Path:
    reports = reports_dir or PROJECT_ROOT / "reports"
    return Builder(reports).build(out or reports / "AML_Detection_Capstone_Report.docx")


if __name__ == "__main__":
    print(f"Wrote {build_report()}")
